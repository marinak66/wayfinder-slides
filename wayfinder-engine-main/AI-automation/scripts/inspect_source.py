"""Quickly inspect the uploaded source .docx: dump paragraphs with their style
names, plus tables, so we can see how 'Slide Header', 'Key message', and
'On slide' are marked up. Read-only."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

SOURCE = Path(__file__).resolve().parents[1] / "inputs" / "For+Slide-Generator+skill.docx"


def main() -> int:
    if not SOURCE.exists():
        print(f"NOT FOUND: {SOURCE}")
        return 1

    doc = Document(str(SOURCE))
    print(f"=== {SOURCE.name} ===")
    print(f"paragraphs: {len(doc.paragraphs)}  tables: {len(doc.tables)}")
    print()

    print("--- PARAGRAPHS (index | style | text) ---")
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style = p.style.name if p.style else "?"
        if not text:
            print(f"{i:4d} | {style:20s} | <empty>")
            continue
        # truncate long lines for readability
        preview = text if len(text) <= 200 else text[:200] + "..."
        print(f"{i:4d} | {style:20s} | {preview}")

    if doc.tables:
        print()
        print(f"--- TABLES: {len(doc.tables)} ---")
        for ti, table in enumerate(doc.tables):
            print(f"\n[Table {ti}] rows={len(table.rows)} cols={len(table.columns)}")
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    cell_text = cell.text.strip().replace("\n", " | ")
                    preview = cell_text if len(cell_text) <= 160 else cell_text[:160] + "..."
                    print(f"  r{ri}c{ci}: {preview}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
