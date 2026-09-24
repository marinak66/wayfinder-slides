"""Verify the generated deck: per slide, count shapes, list a title, tables, and
report the distinct fonts used (should be DT Flow / DT Flow Extrabold)."""

from __future__ import annotations

from collections import Counter
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "outputs" / "projects" / "dynatrace-and-servicenow" / "dynatrace-and-servicenow.pptx"

prs = Presentation(str(DECK))
print(f"Deck: {DECK.name}  |  {len(prs.slides)} slides\n")

all_fonts = Counter()

for i, slide in enumerate(prs.slides):
    n_shapes = len(slide.shapes)
    n_tables = sum(1 for sh in slide.shapes if sh.has_table)
    title = ""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            title = sh.text_frame.text.strip().split("\n")[0][:60]
            break
    print(f"[{i:02d}] shapes={n_shapes:2d} tables={n_tables} | {title}")
    for sh in slide.shapes:
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name:
                        all_fonts[r.font.name] += 1
        if sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    for p in cell.text_frame.paragraphs:
                        for r in p.runs:
                            if r.font.name:
                                all_fonts[r.font.name] += 1

print("\nFonts used across all runs:")
for name, count in all_fonts.most_common():
    print(f"  {name}: {count}")
