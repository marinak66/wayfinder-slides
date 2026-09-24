"""Render the parsed Dynatrace-ServiceNow slide list into a decorated .pptx that
uses the `dt-template/reduced-template.pptx` theme.

Design goals (per user feedback):
- Fonts: headings use "DT Flow Extrabold", body uses "DT Flow" (the theme fonts).
- Never plain text alone: every slide is decorated with frames, cards, bars,
  or path blocks. Tasteful emoji are used as section icons.
- Content is verbatim from the source doc. Layout/format is chosen here.

Per-slide treatment is dispatched by slide number in `render_slide`.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.presentation import Presentation as PresentationT
from pptx.slide import Slide as PptxSlide
from pptx.slide import SlideLayout
from pptx.util import Emu, Inches, Pt

from parse_source import Slide, TableBlock, TextBlock, parse_source

# ---------------------------------------------------------------------------
# Fonts (match the template theme exactly)
# ---------------------------------------------------------------------------
FONT_HEADING = "DT Flow Extrabold"
FONT_BODY = "DT Flow"

# ---------------------------------------------------------------------------
# Brand palette
# ---------------------------------------------------------------------------
DT_NAVY = RGBColor(0x14, 0x22, 0x63)      # deep header navy
DT_PURPLE = RGBColor(0x73, 0x22, 0xB3)    # accent purple
DT_TEAL = RGBColor(0x00, 0xB4, 0xD8)      # accent teal (gradient partner)
DT_BLUE = RGBColor(0x15, 0x6D, 0xF6)      # link/number blue
TABLE_HEADER = RGBColor(0x2E, 0x77, 0xE6)  # lighter table-header blue
HILITE = RGBColor(0xFF, 0xE9, 0xA8)       # marker highlight (soft amber)
SN_GREEN = RGBColor(0x62, 0xB5, 0x4A)     # ServiceNow brand green
SNOW_BLUE = RGBColor(0x29, 0xB5, 0xE8)    # Snowflake brand blue
INK = RGBColor(0x1A, 0x1A, 0x1A)
INK_MUTED = RGBColor(0x55, 0x55, 0x66)
INK_SOFT = RGBColor(0x8A, 0x8A, 0x98)
PANEL_LIGHT = RGBColor(0xF5, 0xF2, 0xFB)
PANEL_BORDER = RGBColor(0xD6, 0xCD, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

GOOD_GREEN = RGBColor(0x1E, 0x7D, 0x3A)
GOOD_FILL = RGBColor(0xEC, 0xF7, 0xEF)
GOOD_BORDER = RGBColor(0xBE, 0xE3, 0xC8)
BAD_RED = RGBColor(0xC0, 0x2B, 0x2B)
BAD_FILL = RGBColor(0xFC, 0xEE, 0xEE)
BAD_BORDER = RGBColor(0xF0, 0xC9, 0xC9)
WARN_AMBER = RGBColor(0xB5, 0x7A, 0x0B)
WARN_FILL = RGBColor(0xFD, 0xF6, 0xE7)
WARN_BORDER = RGBColor(0xEF, 0xDF, 0xB4)

# RACI letter colors
RACI_COLORS = {
    "R": GOOD_GREEN,
    "A": DT_PURPLE,
    "C": WARN_AMBER,
    "I": INK_SOFT,
    "JOINT": DT_PURPLE,
}

# ---------------------------------------------------------------------------
# Geometry (13.333" x 7.5" widescreen)
# ---------------------------------------------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN_L = Inches(0.6)
MARGIN_R = Inches(0.6)
CONTENT_W = SLIDE_W - MARGIN_L - MARGIN_R
ACCENT_TOP = Inches(1.05)
KEY_MSG_TOP = Inches(1.2)
BODY_TOP = Inches(1.9)
BODY_BOTTOM = Inches(7.05)


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def _find_layout(prs: PresentationT, name: str) -> SlideLayout:
    target = name.strip().lower()
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name.strip().lower() == target:
                return layout
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if "standard" in layout.name.lower():
                return layout
    return prs.slide_layouts[0]


def _strip_layout_accent(layout) -> None:
    """Remove the thin decorative accent line the template bakes into the
    'Standard slide' layout between the title and the content — it collided with
    the key message (looked like a strikethrough)."""
    for shp in list(layout.shapes):
        top, h = shp.top, shp.height
        if top is None or h is None:
            continue
        if (shp.name.startswith("Rectangle")
                and h < Inches(0.12)
                and Inches(1.0) <= top <= Inches(1.6)):
            shp._element.getparent().remove(shp._element)


def _clear_slides(prs: PresentationT) -> None:
    part = prs.part
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rId = sld_id.get(qn("r:id"))
        if rId:
            try:
                part.drop_rel(rId)
            except KeyError:
                pass
        sld_id_lst.remove(sld_id)


def _font(run, *, size=None, bold=None, italic=None, color=None, heading=False):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    run.font.name = FONT_HEADING if heading else FONT_BODY


def _no_shadow(shape):
    shape.shadow.inherit = False


def _rounded(slide, left, top, w, h, *, fill, border=None, line_w=0.75, radius=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = border
        shp.line.width = Pt(line_w)
    _no_shadow(shp)
    return shp


def _rect(slide, left, top, w, h, *, fill, border=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if border is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = border
    _no_shadow(shp)
    return shp


def _highlight(run, rgb):
    """Apply a marker-style text highlight to a run (inserted in schema order so
    PowerPoint actually renders it)."""
    rPr = run._r.get_or_add_rPr()
    hl = parse_xml(f'<a:highlight {nsdecls("a")}><a:srgbClr val="{rgb}"/></a:highlight>')
    rPr.insert_element_before(
        hl, "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea",
        "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst",
    )


def _split_label(text):
    """Return (label, rest) when a bullet opens with a short 'Term:' prefix."""
    if ":" not in text:
        return None, text
    label, rest = text.split(":", 1)
    label, rest = label.strip(), rest.strip()
    if 0 < len(label) <= 34 and "." not in label and label[:1].isalnum():
        return label, rest
    return None, text


def _pill(slide, left, top, w, h, text, *, fill, text_color=WHITE, size=13,
          bold=True, heading=True, border=None):
    """Small rounded chip with centered text (used for wordmarks / flow boxes)."""
    shp = _rounded(slide, left, top, w, h, fill=fill, border=border)
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    _font(r, size=size, bold=bold, color=text_color, heading=heading)
    return shp


def _arrow_shape(slide, left, top, w, h, *, color=DT_PURPLE):
    ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, w, h)
    ar.fill.solid()
    ar.fill.fore_color.rgb = color
    ar.line.fill.background()
    _no_shadow(ar)
    return ar


def _flow_row(slide, items, left, top, box_h, *, source_fill):
    """Draw a horizontal sequence of boxes joined by arrows; return end x."""
    x = int(left)
    arrow_w = int(Inches(0.42))
    pad = int(Inches(0.09))
    for i, it in enumerate(items):
        if i > 0:
            _arrow_shape(slide, Emu(x), Emu(int(top) + int(box_h) * 0.30),
                         Emu(arrow_w), Emu(int(int(box_h) * 0.40)), color=DT_PURPLE)
            x += arrow_w + pad
        w = int(Inches(0.125)) * len(it) + int(Inches(0.44))
        if i == 0:
            _pill(slide, Emu(x), top, Emu(w), box_h, it, fill=source_fill,
                  text_color=WHITE, size=13)
        else:
            _pill(slide, Emu(x), top, Emu(w), box_h, it, fill=WHITE,
                  text_color=DT_NAVY, size=13, border=PANEL_BORDER)
        x += w + pad
    return x


def _textbox(slide, left, top, w, h, *, anchor=None, wrap=True):
    tb = slide.shapes.add_textbox(left, top, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    if anchor is not None:
        tf.vertical_anchor = anchor
    return tf


def _para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def _title(slide, text, *, accent=True):
    """Write into the layout title placeholder using DT Flow Extrabold."""
    title_ph = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:
            title_ph = ph
            break
    if title_ph is None:
        tf = _textbox(slide, MARGIN_L, Inches(0.32), CONTENT_W, Inches(0.75))
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        _font(r, size=30, bold=True, color=DT_NAVY, heading=True)
    else:
        tf = title_ph.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = text
        _font(r, bold=True, color=DT_NAVY, heading=True)
    # `accent` retained for signature compatibility; the standalone accent line
    # under the title was removed so the title and key message read as a tight
    # two-line stack (per user feedback).


def _key_message(slide, text):
    if not text:
        return
    # single-line lead statement sitting directly under the title
    tf = _textbox(slide, MARGIN_L, KEY_MSG_TOP, CONTENT_W, Inches(0.5))
    tf.word_wrap = False
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    _font(r, size=15, italic=True, color=INK_MUTED)


def _remove_body_placeholders(slide):
    for ph in list(slide.placeholders):
        idx = ph.placeholder_format.idx
        if idx in (0, 10, 11, 12):
            continue
        ph._element.getparent().remove(ph._element)


def _bullets(tf, items, *, size=14, header_size=None, color=INK,
             header_color=DT_PURPLE, bullet_char="•", space_after=3):
    """Populate a text frame with bullets. Items that look like sub-headers are
    rendered bold, without a bullet."""
    tf.clear()
    first = True
    for raw in items:
        text = raw.strip()
        if not text:
            continue
        p = _para(tf, first=first)
        first = False
        p.alignment = PP_ALIGN.LEFT
        if _is_sub_header(text):
            r = p.add_run()
            r.text = text
            _font(r, size=(header_size or size + 1), bold=True, color=header_color, heading=True)
            p.space_before = Pt(6)
            p.space_after = Pt(2)
        elif text.startswith("#"):
            # Slack channel — bold + marker highlight on the channel token only
            if bullet_char:
                rb = p.add_run(); rb.text = f"{bullet_char}  "
                _font(rb, size=size, color=color)
            parts = text.split(None, 1)
            chan, rest = parts[0], (parts[1] if len(parts) > 1 else "")
            rc = p.add_run(); rc.text = chan
            _font(rc, size=size, bold=True, color=DT_BLUE, heading=True)
            _highlight(rc, HILITE)
            if rest:
                rr = p.add_run(); rr.text = f" {rest}"
                _font(rr, size=size, color=color)
            p.space_after = Pt(space_after)
        else:
            prefix, rest = _split_label(text)
            lead = f"{bullet_char}  " if bullet_char else ""
            if prefix is not None:
                r0 = p.add_run()
                r0.text = f"{lead}{prefix}:"
                _font(r0, size=size, bold=True, color=DT_NAVY, heading=True)
                r1 = p.add_run()
                r1.text = f" {rest}"
                _font(r1, size=size, color=color)
            else:
                r = p.add_run()
                r.text = f"{lead}{text}" if bullet_char else text
                _font(r, size=size, color=color)
            p.space_after = Pt(space_after)


def _is_sub_header(text):
    s = text.strip()
    if len(s) > 90:
        return False
    if s.endswith(":"):
        return True
    low = s.lower()
    return low.startswith((
        "what d1 should say", "what d1 should not say",
        "common misconception", "key takeaways",
    ))


def _panel(slide, *, header, items, left, top, w, h, accent=DT_PURPLE,
           fill=PANEL_LIGHT, border=PANEL_BORDER, emoji="", body_size=13):
    """Rounded panel with a colored header strip and bulleted body."""
    _rounded(slide, left, top, w, h, fill=fill, border=border)
    # header strip
    strip_h = Inches(0.5)
    strip = _rounded(slide, left, top, w, strip_h, fill=accent)
    htf = _textbox(slide, left + Inches(0.2), top + Inches(0.04), w - Inches(0.4), strip_h,
                   anchor=MSO_ANCHOR.MIDDLE)
    hp = htf.paragraphs[0]
    hr = hp.add_run()
    hr.text = f"{emoji}  {header}" if emoji else header
    _font(hr, size=15, bold=True, color=WHITE, heading=True)
    # body
    btf = _textbox(slide, left + Inches(0.22), top + strip_h + Inches(0.1),
                   w - Inches(0.44), h - strip_h - Inches(0.2))
    _bullets(btf, items, size=body_size)


# ---------------------------------------------------------------------------
# Table helper with content-proportional column widths
# ---------------------------------------------------------------------------
def _col_widths(rows, total_w, *, min_w, first_smallest=False):
    ncols = len(rows[0])
    maxlen = [1] * ncols
    for r in rows:
        for i, c in enumerate(r):
            maxlen[i] = max(maxlen[i], len(c))
    if first_smallest:
        # bias the first column down to roughly its own text length
        maxlen[0] = min(maxlen[0], 16)
    tot = sum(maxlen)
    widths = [int(total_w * m / tot) for m in maxlen]
    # enforce minimum
    min_emu = int(min_w)
    deficit = 0
    for i in range(ncols):
        if widths[i] < min_emu:
            deficit += min_emu - widths[i]
            widths[i] = min_emu
    # take the deficit from the widest column
    if deficit:
        widest = max(range(ncols), key=lambda i: widths[i])
        widths[widest] = max(min_emu, widths[widest] - deficit)
    # normalize rounding drift into last column
    drift = int(total_w) - sum(widths)
    widths[-1] += drift
    return widths


def _add_table(slide, block: TableBlock, *, left, top, w, max_h,
               header_size=11, body_size=10, first_smallest=False):
    rows = len(block.rows)
    cols = len(block.rows[0]) if block.rows else 0
    if not rows or not cols:
        return top
    row_h = max(Inches(0.28), min(Inches(0.5), Emu(int(max_h) // rows)))
    shape = slide.shapes.add_table(rows, cols, left, top, w, row_h * rows)
    tbl = shape.table
    # dynamic widths
    widths = _col_widths(block.rows, int(w), min_w=Inches(0.9), first_smallest=first_smallest)
    for ci, wd in enumerate(widths):
        tbl.columns[ci].width = wd
    for ri, row in enumerate(block.rows):
        for ci, cell_text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.07)
            cell.margin_right = Inches(0.07)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.clear()
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = cell_text
            if ri == 0:
                _font(r, size=header_size, bold=True, color=WHITE, heading=True)
                cell.fill.solid()
                cell.fill.fore_color.rgb = TABLE_HEADER
            else:
                _font(r, size=body_size, color=INK)
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if ri % 2 else RGBColor(0xF6, 0xF3, 0xFB)
    return top + row_h * rows


# ---------------------------------------------------------------------------
# Render context + content slide scaffold
# ---------------------------------------------------------------------------
@dataclass
class Ctx:
    prs: PresentationT
    standard: SlideLayout
    cover: SlideLayout


def _content_slide(ctx: Ctx, title, key_msg):
    slide = ctx.prs.slides.add_slide(ctx.standard)
    _remove_body_placeholders(slide)
    _title(slide, title)
    _key_message(slide, key_msg)
    return slide


def _texts(body):
    return [b.text for b in body if isinstance(b, TextBlock)]


def _tables(body):
    return [b for b in body if isinstance(b, TableBlock)]


# ---------------------------------------------------------------------------
# Slide 0 - Cover
# ---------------------------------------------------------------------------
def render_cover(ctx: Ctx):
    slide = ctx.prs.slides.add_slide(ctx.cover)
    _title(slide, "Dynatrace and ServiceNow", accent=False)
    tf = _textbox(slide, MARGIN_L, Inches(4.4), CONTENT_W, Inches(0.7))
    r = tf.paragraphs[0].add_run()
    r.text = "Integration playbook for the D1 field — CoE knowledge session"
    _font(r, size=18, italic=True, color=INK_MUTED)


# ---------------------------------------------------------------------------
# Slide 0 in source (Agenda & Purpose) - framed sections with emoji
# ---------------------------------------------------------------------------
def render_agenda(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    items = _texts(s.body)
    # Expected: Goal / Scale / Naming
    goal = next((t for t in items if t.lower().startswith("goal")), "")
    naming = next((t for t in items if t.lower().startswith("naming")), "")
    scale = next((t for t in items if t not in (goal, naming)), "")

    cards = [
        ("🎯", "Goal", goal.split(":", 1)[1].strip() if ":" in goal else goal, DT_PURPLE),
        ("🌐", "Scale and spread", scale, DT_BLUE),
        ("🏷️", "Naming", naming.split(":", 1)[1].strip() if ":" in naming else naming, DT_TEAL),
    ]
    top = BODY_TOP + Inches(0.05)
    total_h = BODY_BOTTOM - top
    gap = Inches(0.22)
    card_h = Emu(int((int(total_h) - int(gap) * (len(cards) - 1)) // len(cards)))
    for i, (emoji, label, text, accent) in enumerate(cards):
        ctop = Emu(int(top) + (int(card_h) + int(gap)) * i)
        _rounded(slide, MARGIN_L, ctop, CONTENT_W, card_h, fill=PANEL_LIGHT, border=PANEL_BORDER)
        # accent chip on the left
        chip_w = Inches(2.7)
        _rounded(slide, MARGIN_L, ctop, chip_w, card_h, fill=accent)
        ctf = _textbox(slide, MARGIN_L + Inches(0.2), ctop, chip_w - Inches(0.4), card_h,
                       anchor=MSO_ANCHOR.MIDDLE)
        cp = ctf.paragraphs[0]
        cr = cp.add_run()
        cr.text = f"{emoji}  {label}"
        _font(cr, size=18, bold=True, color=WHITE, heading=True)
        # body area
        body_left = MARGIN_L + chip_w + Inches(0.3)
        body_w = CONTENT_W - chip_w - Inches(0.5)
        if label == "Naming":
            # left: verbatim naming text; right: NOW -> ServiceNow / snow -> Snowflake
            text_w = Emu(int(body_w) * 52 // 100)
            btf = _textbox(slide, body_left, ctop, text_w, card_h, anchor=MSO_ANCHOR.MIDDLE)
            br = btf.paragraphs[0].add_run()
            br.text = text
            _font(br, size=13, color=INK)
            # mapping chips on the right
            map_left = Emu(int(body_left) + int(text_w) + int(Inches(0.2)))
            map_right = int(body_left) + int(body_w)
            src_w = Inches(1.15)
            arrow_w = Inches(0.5)
            logo_left = Emu(int(map_left) + int(src_w) + int(arrow_w) + int(Inches(0.24)))
            logo_w = Emu(map_right - int(logo_left))
            row_h = Inches(0.52)
            row_gap = Inches(0.24)
            block_h = int(row_h) * 2 + int(row_gap)
            first_top = int(ctop) + (int(card_h) - block_h) // 2
            mappings = [
                ("NOW", "servicenow", SN_GREEN),
                ("snow", "❄ Snowflake", SNOW_BLUE),
            ]
            for j, (src, logo, fill) in enumerate(mappings):
                rtop = Emu(first_top + (int(row_h) + int(row_gap)) * j)
                _pill(slide, map_left, rtop, src_w, row_h, src, fill=DT_NAVY, size=14)
                _arrow_shape(
                    slide,
                    Emu(int(map_left) + int(src_w) + int(Inches(0.12))),
                    Emu(int(rtop) + int(int(row_h) * 0.28)),
                    arrow_w,
                    Emu(int(int(row_h) * 0.44)),
                    color=DT_PURPLE,
                )
                _pill(slide, logo_left, rtop, logo_w, row_h, logo, fill=fill,
                      size=15, text_color=WHITE)
        else:
            btf = _textbox(slide, body_left, ctop, body_w, card_h, anchor=MSO_ANCHOR.MIDDLE)
            bp = btf.paragraphs[0]
            br = bp.add_run()
            br.text = text
            _font(br, size=14, color=INK)


# ---------------------------------------------------------------------------
# Slide 1 (Why ServiceNow Matters) - stat cards, smaller boxes, bigger font
# ---------------------------------------------------------------------------
def render_stats(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    stats = [t for t in _texts(s.body) if "stats in boxes" not in t.lower()]
    cols, rows = 3, 2
    gap = Inches(0.22)
    grid_left = MARGIN_L
    grid_top = BODY_TOP + Inches(0.1)
    grid_w = CONTENT_W - Inches(0.4)           # slightly narrower (protect logo)
    grid_bottom = Inches(6.45)                 # stop above the corner logo
    grid_h = Emu(int(grid_bottom) - int(grid_top))
    card_w = Emu(int((int(grid_w) - int(gap) * (cols - 1)) // cols))
    card_h = Emu(int((int(grid_h) - int(gap) * (rows - 1)) // rows))
    emojis = ["🏢", "📊", "🔗", "🤝", "📈", "🥇"]
    for idx, text in enumerate(stats[: cols * rows]):
        col, row = idx % cols, idx // cols
        left = Emu(int(grid_left) + (int(card_w) + int(gap)) * col)
        top = Emu(int(grid_top) + (int(card_h) + int(gap)) * row)
        _rounded(slide, left, top, card_w, card_h, fill=WHITE, border=PANEL_BORDER)
        _rect(slide, left, top, card_w, Inches(0.09), fill=DT_PURPLE)
        label, value = _split_stat(text)
        tf = _textbox(slide, left + Inches(0.22), top + Inches(0.2),
                      card_w - Inches(0.44), card_h - Inches(0.3), anchor=MSO_ANCHOR.MIDDLE)
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = f"{emojis[idx % len(emojis)]}  {label}"
        _font(r1, size=14, bold=True, color=DT_NAVY, heading=True)
        if value:
            p2 = tf.add_paragraph()
            p2.space_before = Pt(6)
            r2 = p2.add_run()
            r2.text = value
            _font(r2, size=16, color=INK)


def _split_stat(text):
    for sep in (" - ", " – ", " — "):
        if sep in text:
            a, b = text.split(sep, 1)
            return a.strip(), b.strip()
    if ":" in text:
        a, b = text.split(":", 1)
        return a.strip(), b.strip()
    return text.strip(), ""


# ---------------------------------------------------------------------------
# Slide 2 (Partnership) - bold callout + table
# ---------------------------------------------------------------------------
def render_partnership(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)
    partnership = next((t for t in texts if t.lower().startswith("partnership means")), "")
    table = _tables(s.body)[0] if _tables(s.body) else None

    # callout band with bold "Partnership means:"
    band_top = BODY_TOP
    band_h = Inches(0.7)
    _rounded(slide, MARGIN_L, band_top, CONTENT_W, band_h, fill=PANEL_LIGHT, border=PANEL_BORDER)
    _rect(slide, MARGIN_L, band_top, Inches(0.12), band_h, fill=DT_PURPLE)
    tf = _textbox(slide, MARGIN_L + Inches(0.3), band_top, CONTENT_W - Inches(0.5), band_h,
                  anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    if ":" in partnership:
        label, rest = partnership.split(":", 1)
        r1 = p.add_run(); r1.text = f"🤝  {label}: "
        _font(r1, size=15, bold=True, color=DT_PURPLE, heading=True)
        r2 = p.add_run(); r2.text = rest.strip()
        _font(r2, size=15, color=INK)
    else:
        r = p.add_run(); r.text = partnership
        _font(r, size=15, color=INK)

    if table is not None:
        _add_table(slide, table, left=MARGIN_L, top=band_top + band_h + Inches(0.2),
                   w=CONTENT_W, max_h=Inches(3.6), header_size=12, body_size=11)


# ---------------------------------------------------------------------------
# Slide 3 (ITOM/ITSM Positioning) - concept strip + 3 sentiment boxes
# ---------------------------------------------------------------------------
def render_positioning(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)

    # Group content by the sub-headers present in the source.
    def collect(after_label, stop_labels):
        out, capturing = [], False
        for t in texts:
            low = t.lower()
            if low.startswith(after_label):
                capturing = True
                continue
            if capturing and any(low.startswith(s2) for s2 in stop_labels):
                break
            if capturing:
                out.append(t)
        return out

    say = collect("what d1 should say", ["what d1 should not", "common misconception"])
    notsay = collect("what d1 should not", ["common misconception"])
    misc = collect("common misconception", [])
    concept = next((t for t in texts if "write-path" in t.lower()), "")
    not_confronting = next((t for t in texts if "not confronting" in t.lower()), "")

    # top concept strip with the write-paths drawn as connected boxes
    strip_top = BODY_TOP
    strip_h = Inches(1.62)
    _rounded(slide, MARGIN_L, strip_top, CONTENT_W, strip_h, fill=RGBColor(0xEE, 0xEB, 0xFA),
             border=PANEL_BORDER)
    # header line
    htf = _textbox(slide, MARGIN_L + Inches(0.25), strip_top + Inches(0.08),
                   CONTENT_W - Inches(0.5), Inches(0.34), anchor=MSO_ANCHOR.MIDDLE)
    hr = htf.paragraphs[0].add_run()
    hr.text = "🔀  " + (not_confronting or "ITSM and ITOM are different write-paths")
    _font(hr, size=15, bold=True, color=DT_NAVY, heading=True)

    # split the concept line into a descriptive caption + the arrow flows
    desc, flows_part = concept, ""
    for sep in ("=>", "⇒", "→ ITSM", ":"):
        if sep in concept and ("→" in concept or "->" in concept):
            if sep == "=>" or sep == "⇒":
                desc, flows_part = concept.split(sep, 1)
                break
    if not flows_part and ("→" in concept or "->" in concept):
        # fallback: everything from the first flow token onward
        m = re.search(r"(ITSM\s*(?:→|->).*)$", concept)
        if m:
            desc = concept[: m.start()].rstrip(" -–—=>")
            flows_part = m.group(1)

    flow_rows = []
    for chunk in flows_part.split(";"):
        parts = [p.strip() for p in re.split(r"\s*(?:→|->|➜)\s*", chunk) if p.strip()]
        if parts:
            flow_rows.append(parts)

    if desc.strip():
        dtf = _textbox(slide, MARGIN_L + Inches(0.25), strip_top + Inches(0.44),
                       CONTENT_W - Inches(0.5), Inches(0.3))
        dr = dtf.paragraphs[0].add_run()
        dr.text = desc.strip()
        _font(dr, size=12, italic=True, color=INK_MUTED)

    row_h = Inches(0.42)
    row_gap = Inches(0.12)
    frow_top = strip_top + Inches(0.78)
    for i, row in enumerate(flow_rows[:2]):
        rtop = Emu(int(frow_top) + (int(row_h) + int(row_gap)) * i)
        _flow_row(slide, row, MARGIN_L + Inches(0.3), rtop, row_h,
                  source_fill=(DT_PURPLE if i == 0 else DT_BLUE))

    # three sentiment boxes
    box_top = strip_top + strip_h + Inches(0.2)
    box_h = Emu(int(BODY_BOTTOM) - int(box_top))
    gap = Inches(0.25)
    box_w = Emu(int((int(CONTENT_W) - int(gap) * 2) // 3))
    _panel(slide, header="What D1 should say", items=say, left=MARGIN_L, top=box_top,
           w=box_w, h=box_h, accent=GOOD_GREEN, fill=GOOD_FILL, border=GOOD_BORDER,
           emoji="✅", body_size=11)
    _panel(slide, header="What D1 should NOT say", items=notsay,
           left=Emu(int(MARGIN_L) + int(box_w) + int(gap)), top=box_top,
           w=box_w, h=box_h, accent=BAD_RED, fill=BAD_FILL, border=BAD_BORDER,
           emoji="🚫", body_size=11)
    _panel(slide, header="Common misconception", items=misc,
           left=Emu(int(MARGIN_L) + (int(box_w) + int(gap)) * 2), top=box_top,
           w=box_w, h=box_h, accent=WARN_AMBER, fill=WARN_FILL, border=WARN_BORDER,
           emoji="⚠️", body_size=9.5)


# ---------------------------------------------------------------------------
# Slide 4 (Reference Architecture) - graphic slot left, text panel right
# ---------------------------------------------------------------------------
def render_reference_arch(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = [t for t in _texts(s.body)
             if t.strip().lower() not in ("reference architecture graphic", "bullet points:")]
    # flatten multi-line bullet block from source
    items = []
    for t in texts:
        for line in t.split("\n"):
            line = line.strip()
            if line and line.lower() not in ("bullet points:",):
                items.append(line)

    # left: graphic placeholder frame (user inserts image later)
    gap = Inches(0.3)
    left_w = Inches(7.2)
    left = MARGIN_L
    top = BODY_TOP
    h = Emu(int(BODY_BOTTOM) - int(top))
    ph = _rounded(slide, left, top, left_w, h, fill=RGBColor(0xF2, 0xF2, 0xF6),
                  border=PANEL_BORDER)
    ptf = _textbox(slide, left, top, left_w, h, anchor=MSO_ANCHOR.MIDDLE)
    pp = ptf.paragraphs[0]; pp.alignment = PP_ALIGN.CENTER
    pr = pp.add_run(); pr.text = "🖼️  Reference architecture graphic\n(insert here)"
    _font(pr, size=15, italic=True, color=INK_SOFT)

    # right: compact text panel
    right_left = Emu(int(left) + int(left_w) + int(gap))
    right_w = Emu(int(CONTENT_W) - int(left_w) - int(gap))
    _panel(slide, header="Where RCA & remediation fire", items=items,
           left=right_left, top=top, w=right_w, h=h, accent=DT_PURPLE, emoji="🧭",
           body_size=12)


# ---------------------------------------------------------------------------
# Slide 5 (RACI) - one combined table + colored callout at the bottom
# ---------------------------------------------------------------------------
def render_raci(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    # pair each section label with the table that follows it
    sections = []  # (label, TableBlock)
    pending_label = None
    callout = ""
    for b in s.body:
        if isinstance(b, TextBlock):
            low = b.text.lower()
            if low.startswith("callout"):
                callout = b.text.split(":", 1)[1].strip() if ":" in b.text else b.text
            elif b.text.strip().rstrip(":") .lower() != "raci tables, grouped per section":
                pending_label = b.text.rstrip(":").strip()
        elif isinstance(b, TableBlock):
            sections.append((pending_label or "", b))
            pending_label = None

    # build one unified table: col-header row + (section header + data rows)*
    total_rows = 1
    for _, tbl in sections:
        total_rows += 1 + len(tbl.data_rows)
    cols = 3

    top = BODY_TOP
    callout_h = Inches(0.85) if callout else Inches(0)
    table_bottom = Emu(int(BODY_BOTTOM) - int(callout_h) - (int(Inches(0.2)) if callout else 0))
    avail_h = int(table_bottom) - int(top)
    row_h = Emu(max(int(Inches(0.24)), avail_h // total_rows))

    shape = slide.shapes.add_table(total_rows, cols, MARGIN_L, top, CONTENT_W, Emu(int(row_h) * total_rows))
    tbl = shape.table
    tbl.columns[0].width = Inches(8.0)
    tbl.columns[1].width = Inches(2.06)
    tbl.columns[2].width = Inches(2.07)
    tbl.first_row = False

    def set_cell(r, c, text, *, size, color, bold=False, fill=None, align=None, heading=False):
        cell = tbl.cell(r, c)
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.06)
        cell.margin_top = Inches(0.01)
        cell.margin_bottom = Inches(0.01)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf = cell.text_frame
        tf.word_wrap = True
        tf.clear()
        p = tf.paragraphs[0]
        if align:
            p.alignment = align
        run = p.add_run()
        run.text = text
        _font(run, size=size, bold=bold, color=color, heading=heading)
        if fill is not None:
            cell.fill.solid()
            cell.fill.fore_color.rgb = fill

    # column header row
    for c, label in enumerate(["Area", "Dynatrace", "ServiceNow"]):
        set_cell(0, c, label, size=11, color=WHITE, bold=True, fill=TABLE_HEADER,
                 align=None if c == 0 else PP_ALIGN.CENTER, heading=True)

    r = 1
    alt = 0
    for label, sec in sections:
        # section header row spans all 3 columns
        merged = tbl.cell(r, 0)
        merged.merge(tbl.cell(r, 2))
        set_cell(r, 0, f"▸  {label}", size=10.5, color=WHITE, bold=True,
                 fill=DT_PURPLE, heading=True)
        r += 1
        for drow in sec.data_rows:
            area = drow[0]
            dt_v = drow[1] if len(drow) > 1 else ""
            sn_v = drow[2] if len(drow) > 2 else ""
            row_fill = WHITE if alt % 2 else RGBColor(0xF6, 0xF3, 0xFB)
            alt += 1
            set_cell(r, 0, area, size=9.5, color=INK, fill=row_fill)
            set_cell(r, 1, dt_v, size=9.5, color=RACI_COLORS.get(dt_v.upper(), INK),
                     bold=True, fill=row_fill, align=PP_ALIGN.CENTER, heading=True)
            set_cell(r, 2, sn_v, size=9.5, color=RACI_COLORS.get(sn_v.upper(), INK),
                     bold=True, fill=row_fill, align=PP_ALIGN.CENTER, heading=True)
            r += 1

    # colored callout at the bottom
    if callout:
        ctop = Emu(int(table_bottom) + int(Inches(0.2)))
        _rounded(slide, MARGIN_L, ctop, CONTENT_W, callout_h, fill=WARN_FILL, border=WARN_BORDER)
        _rect(slide, MARGIN_L, ctop, Inches(0.12), callout_h, fill=WARN_AMBER)
        tf = _textbox(slide, MARGIN_L + Inches(0.3), ctop, CONTENT_W - Inches(0.5), callout_h,
                      anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r1 = p.add_run(); r1.text = "⚠️  "
        _font(r1, size=12, bold=True, color=WARN_AMBER, heading=True)
        r2 = p.add_run(); r2.text = callout
        _font(r2, size=12, color=INK)


# ---------------------------------------------------------------------------
# Slides 6 & 7 (TODAY / FUTURE) - framed placeholder note
# ---------------------------------------------------------------------------
def render_placeholder(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    top = Inches(2.7)
    h = Inches(2.2)
    _rounded(slide, Inches(2.2), top, Inches(8.9), h, fill=PANEL_LIGHT, border=PANEL_BORDER)
    tf = _textbox(slide, Inches(2.2), top, Inches(8.9), h, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "🚧  To be populated by Alexander Mohr"
    _font(r, size=22, italic=True, color=INK_SOFT, heading=True)


# ---------------------------------------------------------------------------
# Slide 8 (Roadmap Disclosure Rules) - path blocks
# ---------------------------------------------------------------------------
def render_roadmap_path(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)
    rule = next((t for t in texts if t.lower().startswith("if a customer asks")), "")
    contact = next((t for t in texts if "coe contact" in t.lower() or "@dynatrace" in t.lower()), "")

    # top rule band
    band_top = BODY_TOP
    band_h = Inches(0.9)
    _rounded(slide, MARGIN_L, band_top, CONTENT_W, band_h, fill=RGBColor(0xEE, 0xEB, 0xFA),
             border=PANEL_BORDER)
    tf = _textbox(slide, MARGIN_L + Inches(0.25), band_top, CONTENT_W - Inches(0.5), band_h,
                  anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    r0 = p.add_run(); r0.text = "🗣️  Rule:  "
    _font(r0, size=14, bold=True, color=DT_PURPLE, heading=True)
    r1 = p.add_run(); r1.text = rule
    _font(r1, size=14, color=INK)

    # escalation path as chevron blocks
    steps = [
        ("👤", "D1 field", "works with IAT / AE / SE"),
        ("🤝", "Account SE", "properly engaged"),
        ("🧭", "Principal SE", "leads the roadmap call"),
    ]
    path_top = band_top + band_h + Inches(0.45)
    path_h = Inches(1.9)
    gap = Inches(0.25)
    step_w = Emu(int((int(CONTENT_W) - int(gap) * (len(steps) - 1)) // len(steps)))
    label_tf = _textbox(slide, MARGIN_L, band_top + band_h + Inches(0.08), CONTENT_W, Inches(0.35))
    lr = label_tf.paragraphs[0].add_run(); lr.text = "Escalation order"
    _font(lr, size=13, bold=True, color=DT_NAVY, heading=True)
    for i, (emoji, title, sub) in enumerate(steps):
        left = Emu(int(MARGIN_L) + (int(step_w) + int(gap)) * i)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.PENTAGON if i < len(steps) - 1 else MSO_SHAPE.ROUNDED_RECTANGLE,
            left, path_top, step_w, path_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = [DT_PURPLE, DT_BLUE, DT_TEAL][i]
        shape.line.fill.background()
        _no_shadow(shape)
        stf = shape.text_frame
        stf.word_wrap = True
        stf.vertical_anchor = MSO_ANCHOR.MIDDLE
        sp = stf.paragraphs[0]; sp.alignment = PP_ALIGN.CENTER
        sr = sp.add_run(); sr.text = f"{emoji}\n{title}"
        _font(sr, size=16, bold=True, color=WHITE, heading=True)
        sp2 = stf.add_paragraph(); sp2.alignment = PP_ALIGN.CENTER
        sr2 = sp2.add_run(); sr2.text = sub
        _font(sr2, size=11, color=WHITE)

    # contact block
    ctop = path_top + path_h + Inches(0.35)
    _rounded(slide, MARGIN_L, ctop, CONTENT_W, Inches(0.7), fill=PANEL_LIGHT, border=PANEL_BORDER)
    _rect(slide, MARGIN_L, ctop, Inches(0.12), Inches(0.7), fill=DT_TEAL)
    ctf = _textbox(slide, MARGIN_L + Inches(0.3), ctop, CONTENT_W - Inches(0.5), Inches(0.7),
                   anchor=MSO_ANCHOR.MIDDLE)
    cp = ctf.paragraphs[0]
    cr = cp.add_run(); cr.text = "📧  "
    _font(cr, size=13, bold=True, color=DT_TEAL, heading=True)
    if ":" in contact:
        clabel, cval = contact.split(":", 1)
        cr1 = cp.add_run(); cr1.text = f"{clabel.strip()}:  "
        _font(cr1, size=13, color=INK)
        cr2 = cp.add_run(); cr2.text = cval.strip()
        _font(cr2, size=13, bold=True, color=DT_NAVY, heading=True)
        _highlight(cr2, HILITE)
    else:
        cr2 = cp.add_run(); cr2.text = contact
        _font(cr2, size=13, bold=True, color=DT_NAVY, heading=True)
        _highlight(cr2, HILITE)


# ---------------------------------------------------------------------------
# Slide 9 (Workarounds) - two panels, bigger font
# ---------------------------------------------------------------------------
def render_workarounds(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)
    left_header, right_header = "Migration", "CoE workarounds"
    left_items, right_items = [], []
    current = None
    for t in texts:
        low = t.lower()
        if low.startswith("2 ") and "box" in low:
            continue
        if "left box" in low:
            current = left_items
            for sep in (" - ", " – ", " — "):
                if sep in t:
                    left_header = t.split(sep, 1)[1].strip()
                    break
            continue
        if "right box" in low:
            current = right_items
            for sep in (" - ", " – ", " — "):
                if sep in t:
                    right_header = t.split(sep, 1)[1].split("(")[0].strip()
                    break
            continue
        (current if current is not None else left_items).append(t)

    top = BODY_TOP
    h = Emu(int(BODY_BOTTOM) - int(top))
    gap = Inches(0.3)
    panel_w = Emu(int((int(CONTENT_W) - int(gap)) // 2))
    _panel(slide, header=left_header, items=left_items, left=MARGIN_L, top=top,
           w=panel_w, h=h, accent=DT_PURPLE, emoji="🧭", body_size=14)
    _panel(slide, header=right_header, items=right_items,
           left=Emu(int(MARGIN_L) + int(panel_w) + int(gap)), top=top,
           w=panel_w, h=h, accent=DT_BLUE, emoji="🛠️", body_size=14)


# ---------------------------------------------------------------------------
# Slide 10 (DEMO) - fitted table + bottom callout (merged with notes)
# ---------------------------------------------------------------------------
def render_demo(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    table = _tables(s.body)[0]
    # notes after the table
    seen_table = False
    notes = []
    for b in s.body:
        if isinstance(b, TableBlock):
            seen_table = True
            continue
        if seen_table and isinstance(b, TextBlock):
            notes.append(b.text)

    top = BODY_TOP
    callout_h = Inches(1.0)
    table_bottom = Emu(int(BODY_BOTTOM) - int(callout_h) - int(Inches(0.2)))
    _add_table(slide, table, left=MARGIN_L, top=top, w=CONTENT_W,
               max_h=Emu(int(table_bottom) - int(top)),
               header_size=9.5, body_size=8.5, first_smallest=True)

    # callout with pipeline + template rule
    ctop = Emu(int(table_bottom) + int(Inches(0.2)))
    _rounded(slide, MARGIN_L, ctop, CONTENT_W, callout_h, fill=RGBColor(0xEE, 0xEB, 0xFA),
             border=PANEL_BORDER)
    _rect(slide, MARGIN_L, ctop, Inches(0.12), callout_h, fill=DT_PURPLE)
    tf = _textbox(slide, MARGIN_L + Inches(0.3), ctop, CONTENT_W - Inches(0.5), callout_h,
                  anchor=MSO_ANCHOR.MIDDLE)
    first = True
    for n in notes:
        p = _para(tf, first=first); first = False
        r = p.add_run()
        r.text = f"🔧  {n}" if n.lower().startswith("template") else f"➡️  {n}"
        _font(r, size=11, color=INK)
        p.space_after = Pt(2)


# ---------------------------------------------------------------------------
# Slide 11 (Wrap-up) - two-column takeaways + help channels
# ---------------------------------------------------------------------------
def render_wrapup(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)
    takeaways = [t for t in texts if _is_takeaway(t)]
    channels = [t for t in texts if t.strip().startswith("#")]
    feedback = next((t for t in texts if t.lower().startswith("feedback loop")), "")
    library = next((t for t in texts if "already written up" in t.lower()), "")

    top = BODY_TOP
    h = Emu(int(BODY_BOTTOM) - int(top))
    gap = Inches(0.3)
    left_w = Inches(6.9)
    # left: key takeaways as numbered list panel
    _panel(slide, header="Key takeaways", items=[f"{i+1}. {t}" for i, t in enumerate(takeaways)],
           left=MARGIN_L, top=top, w=left_w, h=h, accent=DT_PURPLE, emoji="📌", body_size=12)

    # right: help + feedback
    right_left = Emu(int(MARGIN_L) + int(left_w) + int(gap))
    right_w = Emu(int(CONTENT_W) - int(left_w) - int(gap))
    right_items = []
    if library:
        right_items.append(library)
    right_items.extend(channels)
    right_items.append("Don't cross-post the same question to both channels")
    if feedback:
        right_items.append(feedback)
    _panel(slide, header="Where to get help", items=right_items, left=right_left, top=top,
           w=right_w, h=h, accent=DT_BLUE, emoji="💬", body_size=11.5)


def _is_takeaway(t):
    low = t.lower()
    if t.strip().startswith("#"):
        return False
    if low.startswith(("key takeaways", "everything above", "requests", "don't cross", "feedback loop")):
        return False
    # takeaways are the four substantive summary lines
    return len(t) > 80


# ---------------------------------------------------------------------------
# Slide 12a (Appendix FAQ) - cheat-sheet card
# ---------------------------------------------------------------------------
def render_faq(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    texts = _texts(s.body)
    questions = [t for t in texts if t.strip().endswith("?")]
    note = next((t for t in texts if t.lower().startswith("note")), "")
    link = next((t for t in texts if "link" in t.lower()), "")

    top = BODY_TOP
    # link placeholder chip
    _rounded(slide, MARGIN_L, top, CONTENT_W, Inches(0.7), fill=PANEL_LIGHT, border=PANEL_BORDER)
    _rect(slide, MARGIN_L, top, Inches(0.12), Inches(0.7), fill=DT_TEAL)
    ltf = _textbox(slide, MARGIN_L + Inches(0.3), top, CONTENT_W - Inches(0.5), Inches(0.7),
                   anchor=MSO_ANCHOR.MIDDLE)
    lp = ltf.paragraphs[0]
    lr = lp.add_run(); lr.text = "🔗  " + (link or "Link to the cheat sheet (insert)")
    _font(lr, size=13, italic=True, color=INK_MUTED)

    # question cards
    q_top = top + Inches(0.95)
    q_h = Inches(0.85)
    q_gap = Inches(0.2)
    for i, q in enumerate(questions):
        qtop = Emu(int(q_top) + (int(q_h) + int(q_gap)) * i)
        _rounded(slide, MARGIN_L, qtop, CONTENT_W, q_h, fill=WHITE, border=PANEL_BORDER)
        _rounded(slide, MARGIN_L, qtop, Inches(0.7), q_h, fill=DT_PURPLE)
        etf = _textbox(slide, MARGIN_L, qtop, Inches(0.7), q_h, anchor=MSO_ANCHOR.MIDDLE)
        ep = etf.paragraphs[0]; ep.alignment = PP_ALIGN.CENTER
        er = ep.add_run(); er.text = "❓"
        _font(er, size=18, color=WHITE)
        qtf = _textbox(slide, MARGIN_L + Inches(0.9), qtop, CONTENT_W - Inches(1.1), q_h,
                       anchor=MSO_ANCHOR.MIDDLE)
        qr = qtf.paragraphs[0].add_run(); qr.text = q
        _font(qr, size=14, color=INK)

    # note footer
    if note:
        ntf = _textbox(slide, MARGIN_L, Inches(6.6), CONTENT_W, Inches(0.5))
        nr = ntf.paragraphs[0].add_run(); nr.text = "📄  " + note
        _font(nr, size=12, italic=True, color=INK_SOFT)


# ---------------------------------------------------------------------------
# Slide 12b (QA Live) - decorative Q&A
# ---------------------------------------------------------------------------
def render_qa(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    top = Inches(2.6)
    h = Inches(2.4)
    _rounded(slide, Inches(3.4), top, Inches(6.5), h, fill=DT_NAVY)
    tf = _textbox(slide, Inches(3.4), top, Inches(6.5), h, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "💬  Q & A"
    _font(r, size=40, bold=True, color=WHITE, heading=True)
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(8)
    r2 = p2.add_run(); r2.text = "Pick questions from the box — unanswered ones get a follow-up to D1"
    _font(r2, size=14, italic=True, color=RGBColor(0xD8, 0xD2, 0xEC))


# ---------------------------------------------------------------------------
# Fallback generic bullets
# ---------------------------------------------------------------------------
def render_generic(ctx: Ctx, s: Slide):
    slide = _content_slide(ctx, s.title, s.key_message)
    items = _texts(s.body)
    # frame the bullets in a light panel so it isn't bare text
    top = BODY_TOP
    h = Emu(int(BODY_BOTTOM) - int(top))
    _rounded(slide, MARGIN_L, top, CONTENT_W, h, fill=PANEL_LIGHT, border=PANEL_BORDER)
    tf = _textbox(slide, MARGIN_L + Inches(0.3), top + Inches(0.2),
                  CONTENT_W - Inches(0.6), h - Inches(0.4))
    _bullets(tf, items, size=14)


# ---------------------------------------------------------------------------
# Dispatcher
# ---------------------------------------------------------------------------
def render_slide(ctx: Ctx, s: Slide):
    num = s.number
    title_low = s.title.lower()
    if num == "0":
        render_agenda(ctx, s)
    elif num == "1":
        render_stats(ctx, s)
    elif num == "2":
        render_partnership(ctx, s)
    elif num == "3":
        render_positioning(ctx, s)
    elif num == "4":
        render_reference_arch(ctx, s)
    elif num == "5":
        render_raci(ctx, s)
    elif num in ("6", "7"):
        render_placeholder(ctx, s)
    elif num == "8":
        render_roadmap_path(ctx, s)
    elif num == "9":
        render_workarounds(ctx, s)
    elif num == "10":
        render_demo(ctx, s)
    elif num == "11":
        render_wrapup(ctx, s)
    elif num == "12" and "appendix" in title_low:
        render_faq(ctx, s)
    elif num == "12":
        render_qa(ctx, s)
    else:
        render_generic(ctx, s)


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------
def build_deck(source_docx: Path, template_pptx: Path, output_pptx: Path):
    slides = parse_source(source_docx)
    prs = Presentation(str(template_pptx))
    _clear_slides(prs)
    ctx = Ctx(
        prs=prs,
        standard=_find_layout(prs, "Standard slide"),
        cover=_find_layout(prs, "Cover slide"),
    )
    _strip_layout_accent(ctx.standard)
    render_cover(ctx)
    for s in slides:
        render_slide(ctx, s)
    output_pptx.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_pptx))
    print(f"Wrote {output_pptx}  ({len(prs.slides)} slides)")


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1]
    src = root / "inputs" / "For+Slide-Generator+skill.docx"
    tpl = root / "dt-template" / "reduced-template.pptx"
    out = root / "outputs" / "projects" / "dynatrace-and-servicenow" / "dynatrace-and-servicenow.pptx"
    build_deck(src, tpl, out)
